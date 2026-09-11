"""
app/ai/text_extraction.py

Single source of truth for extracting plain text from an uploaded
document (PDF/DOCX/XLSX). Used in two places:

1. documents.py, at upload time (submit_document) — extracts text from
   the raw uploaded bytes and stores it on Document.extracted_text, so
   it's available for display immediately without re-reading the file.
2. service.py, in the AI assist pipeline — prefers doc.extracted_text
   (already computed at upload) over re-extracting, falling back to a
   fresh extraction only for documents uploaded before this existed.

Never raises on a bad/corrupt file — returns an "[Error extracting
text: ...]" string instead, so a malformed upload never breaks the
submit_document request itself.
"""

from __future__ import annotations

import io

import openpyxl
import pypdf
from docx import Document as DocxDocument

from app.models import Document

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def extract_file_text(file_bytes: bytes, content_type: str, filename: str) -> str:
    """Extract plain text from raw file bytes, given its content type and
    filename. This is the function to call at upload time, before the
    file is necessarily saved to disk."""
    try:
        if content_type == PDF_MIME or filename.endswith(".pdf"):
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            return "\n".join(page.extract_text() for page in reader.pages if page.extract_text())

        elif content_type == DOCX_MIME or filename.endswith(".docx"):
            doc = DocxDocument(io.BytesIO(file_bytes))
            return "\n".join(para.text for para in doc.paragraphs if para.text)

        elif content_type == XLSX_MIME or filename.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            text_acc = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row in ws.iter_rows(values_only=True):
                    row_str = " ".join(str(cell) for cell in row if cell is not None)
                    if row_str:
                        text_acc.append(row_str)
            return "\n".join(text_acc)

    except Exception as e:
        return f"[Error extracting text: {str(e)}]"

    return file_bytes.decode("utf-8", errors="ignore")


def extract_text(doc: Document) -> str:
    """Back-compat path for callers that only have a Document row, not
    raw bytes — reads the file from disk via doc.file_path and extracts.
    Prefer using doc.extracted_text directly wherever it's already been
    populated (i.e. any document uploaded after this feature landed);
    this function is now mainly a fallback for older documents uploaded
    before extraction was wired into the upload flow."""
    with open(doc.file_path, "rb") as f:
        file_bytes = f.read()
    return extract_file_text(file_bytes, doc.content_type, doc.filename)